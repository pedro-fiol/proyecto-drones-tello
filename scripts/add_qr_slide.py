"""Add QR-code slide to TelloInterceptor.pptx pointing at the GitHub repo."""

from pathlib import Path

import qrcode
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

REPO_URL = "https://github.com/Aszent47/proyecto-drones-tello"
PROJECT_ROOT = Path(r"C:/AEROS UPC/4A/PD/proyecto-drones-tello")
PPTX_PATH = PROJECT_ROOT / "TelloInterceptor.pptx"
QR_PATH = PROJECT_ROOT / "assets" / "qr_repo_access.png"


def generate_qr(url: str, out_path: Path) -> None:
    """Render high-contrast QR PNG sized for slide-deck embedding."""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    qr = qrcode.QRCode(
        version=None,
        error_correction=qrcode.constants.ERROR_CORRECT_H,
        box_size=20,
        border=2,
    )
    qr.add_data(url)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    img.save(out_path)


def add_qr_slide(pptx_path: Path, qr_path: Path, repo_url: str) -> None:
    """Append a new slide showing the QR code and repository URL."""
    prs = Presentation(str(pptx_path))
    slide_w = prs.slide_width
    slide_h = prs.slide_height

    blank_layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]
    slide = prs.slides.add_slide(blank_layout)

    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, slide_w, slide_h)
    bg.line.fill.background()
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(0x10, 0x12, 0x1A)

    title_box = slide.shapes.add_textbox(Emu(0), Emu(700000), slide_w, Emu(800000))
    title_tf = title_box.text_frame
    title_tf.margin_left = Emu(0)
    title_tf.margin_right = Emu(0)
    p = title_tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "Repositorio"
    run.font.size = Pt(44)
    run.font.bold = True
    run.font.name = "Calibri"
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    qr_size = Emu(4500000)
    qr_left = Emu(int((slide_w - qr_size) / 2))
    qr_top = Emu(1900000)
    qr_bg_pad = Emu(180000)
    qr_bg = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        qr_left - qr_bg_pad,
        qr_top - qr_bg_pad,
        qr_size + qr_bg_pad * 2,
        qr_size + qr_bg_pad * 2,
    )
    qr_bg.line.fill.background()
    qr_bg.fill.solid()
    qr_bg.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    slide.shapes.add_picture(str(qr_path), qr_left, qr_top, qr_size, qr_size)

    url_top = qr_top + qr_size + qr_bg_pad + Emu(250000)
    url_box = slide.shapes.add_textbox(Emu(0), url_top, slide_w, Emu(500000))
    url_tf = url_box.text_frame
    url_tf.margin_left = Emu(0)
    url_tf.margin_right = Emu(0)
    p = url_tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = repo_url
    run.font.size = Pt(20)
    run.font.name = "Consolas"
    run.font.color.rgb = RGBColor(0xCA, 0xDC, 0xFC)

    cap_box = slide.shapes.add_textbox(Emu(0), url_top + Emu(550000), slide_w, Emu(700000))
    cap_tf = cap_box.text_frame
    cap_tf.margin_left = Emu(0)
    cap_tf.margin_right = Emu(0)
    p = cap_tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "Escanea para solicitar acceso  ·  invitación manual como colaborador"
    run.font.size = Pt(16)
    run.font.italic = True
    run.font.name = "Calibri"
    run.font.color.rgb = RGBColor(0x97, 0xA8, 0xC8)

    prs.save(str(pptx_path))


if __name__ == "__main__":
    generate_qr(REPO_URL, QR_PATH)
    add_qr_slide(PPTX_PATH, QR_PATH, REPO_URL)
    print(f"QR saved: {QR_PATH}")
    print(f"Slide appended to: {PPTX_PATH}")
